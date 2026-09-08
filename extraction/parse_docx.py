"""Stage 1 parser: DOCX -> raw block list, via `python-docx`."""
from __future__ import annotations

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

HEADING_STYLES = {"Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"}


def parse_docx(path) -> list[dict]:
    """Parse one `.docx` into an ordered list of raw blocks, in document order.

    `python-docx`'s `Document.paragraphs` and `Document.tables` are separate
    flat collections that lose the original interleaving -- walking
    `document.element.body` directly instead preserves heading/paragraph/table
    order, which matters here because a table's nearest preceding heading
    becomes its `section` in the final chunk.

    Each block is one of:
      {"type": "heading", "text": str}
      {"type": "text", "text": str}
      {"type": "table", "rows": list[list[str]]}

    docx has no native page boundaries pre-render, so blocks carry no page
    number -- see `common.py`'s chunk schema docstring and
    `docs/limitations.md` for why that's a deliberate omission, not a gap.
    """
    document = docx.Document(path)
    blocks: list[dict] = []

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            block_type = "heading" if paragraph.style.name in HEADING_STYLES else "text"
            blocks.append({"type": block_type, "text": text})

        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if any(any(cell for cell in row) for row in rows):
                blocks.append({"type": "table", "rows": rows})

    return blocks
